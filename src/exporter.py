import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any

def export_to_excel(detailed_data: List[Dict[str, Any]], summary_data: List[Dict[str, Any]], output_filepath: str):
    """
    Exporta el inventario a un archivo Excel.
    """
    df_summary = pd.DataFrame(summary_data)
    df_detailed = pd.DataFrame(detailed_data)

    with pd.ExcelWriter(output_filepath, engine='openpyxl') as writer:
        df_summary.to_excel(writer, sheet_name='Matriz_Resumen_Pozos', index=False)
        df_detailed.to_excel(writer, sheet_name='Detalle_General_Archivos', index=False)
        
        workbook = writer.book
        
        # Estilos visuales
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Azul Petróleo
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        # 1. Aplicar estilo a la hoja "Matriz_Resumen_Pozos"
        ws_summary = writer.sheets['Matriz_Resumen_Pozos']
        
        # Formatear encabezados
        for col in ws_summary.iter_cols(min_row=1, max_row=1):
            for cell in col:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
        # Formatear celdas de contenido (salto de línea y alineación superior)
        for row in ws_summary.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Ajustar ancho de columnas
        for col in ws_summary.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_summary.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        # 2. Aplicar estilo a la hoja "Detalle_General_Archivos"
        ws_detailed = writer.sheets['Detalle_General_Archivos']
        for col in ws_detailed.iter_cols(min_row=1, max_row=1):
            for cell in col:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                
        for col in ws_detailed.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_detailed.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    print(f"\n[Éxito] Inventario exportado correctamente en: {output_filepath}")